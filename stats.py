import io
import os
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import db

TZ = ZoneInfo(os.getenv("TZ", "Europe/Moscow"))


def channel_base(w):
    ref = (w["channel_ref"] or "") if w else ""
    if ref.startswith("@"):
        return f"https://t.me/{ref[1:]}"
    if w and w["channel_id"]:
        return f"https://t.me/c/{abs(w['channel_id']) - 10 ** 12}"
    return ""


def post_link(w, unit):
    if not w or not unit:
        return ""
    first = db.unit_first(w["id"], unit)
    base = channel_base(w)
    return f"{base}/{first}" if first and base else ""


def _dt(ts):
    return datetime.fromtimestamp(ts, TZ).replace(tzinfo=None) if ts else ""


def _sheet(ws, header, rows):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_xlsx() -> bytes:
    workers = {w["id"]: w for w in db.list_workers()}
    log = db.conn.execute("SELECT * FROM log ORDER BY ts DESC").fetchall()

    def bname(r):
        w = workers.get(r["worker_id"])
        return r["worker_name"] or (w["name"] if w else f"бот {r['worker_id']}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал"
    _sheet(ws, ["Дата и время", "Бот", "Чат", "ID чата", "Пост", "Ссылка на пост", "Статус", "Ошибка"],
           [[_dt(r["ts"]), bname(r), r["chat_title"], r["chat_id"], r["unit"],
             post_link(workers.get(r["worker_id"]), r["unit"]), r["status"], r["error"]]
            for r in log])

    ws = wb.create_sheet("По ботам")
    rows = db.conn.execute(
        "SELECT worker_id, MAX(worker_name) worker_name, SUM(status='ok') ok, SUM(status!='ok') err, "
        "COUNT(DISTINCT chat_id) chats, MAX(CASE WHEN status='ok' THEN ts END) last "
        "FROM log GROUP BY worker_id").fetchall()
    _sheet(ws, ["Бот", "Успешно", "Ошибки", "Чатов в журнале", "Последний пост"],
           [[bname(r), r["ok"], r["err"], r["chats"], _dt(r["last"])] for r in rows])

    ws = wb.create_sheet("По чатам")
    rows = db.conn.execute(
        "SELECT worker_id, MAX(worker_name) wn, chat_id, MAX(chat_title) t, SUM(status='ok') ok, "
        "SUM(status!='ok') err, MAX(CASE WHEN status='ok' THEN ts END) last "
        "FROM log GROUP BY worker_id, chat_id").fetchall()
    _sheet(ws, ["Бот", "Чат", "ID чата", "Успешно", "Ошибки", "Последний пост"],
           [[bname({"worker_name": r["wn"], "worker_id": r["worker_id"]}), r["t"], r["chat_id"],
             r["ok"], r["err"], _dt(r["last"])] for r in rows])

    ws = wb.create_sheet("По постам")
    rows = db.conn.execute(
        "SELECT worker_id, MAX(worker_name) wn, unit, SUM(status='ok') ok, "
        "COUNT(DISTINCT CASE WHEN status='ok' THEN chat_id END) chats "
        "FROM log WHERE unit!='' GROUP BY worker_id, unit ORDER BY ok DESC").fetchall()
    _sheet(ws, ["Бот", "Пост", "Ссылка", "Отправок", "Чатов"],
           [[bname({"worker_name": r["wn"], "worker_id": r["worker_id"]}), r["unit"],
             post_link(workers.get(r["worker_id"]), r["unit"]), r["ok"], r["chats"]] for r in rows])

    ws = wb.create_sheet("По дням")
    days = {}
    for r in log:
        d = datetime.fromtimestamp(r["ts"], TZ).date()
        ok, err = days.get(d, (0, 0))
        days[d] = (ok + (r["status"] == "ok"), err + (r["status"] != "ok"))
    _sheet(ws, ["Дата", "Успешно", "Ошибки"], [[d, *v] for d, v in sorted(days.items(), reverse=True)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
